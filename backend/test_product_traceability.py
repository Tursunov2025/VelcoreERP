"""Disposable acceptance for unit passport generation and QR resolution."""
from __future__ import annotations

import os
import sys
import tempfile
from datetime import datetime
from pathlib import Path

DB = tempfile.NamedTemporaryFile(suffix=".db", delete=False).name
ROOT = Path(tempfile.mkdtemp(prefix="velcore-traceability-"))
os.environ.update({
    "DATABASE_URL": f"sqlite:///{DB.replace(chr(92), '/')}",
    "DB_PATH": DB,
    "DATA_ROOT": str(ROOT),
    "DATABASE_GUARD": "false",
    "ENVIRONMENT": "test",
    "JWT_SECRET_KEY": "traceability-test-secret",
    "PUBLIC_FRONTEND_URL": "https://trace.test",
})
sys.path.insert(0, str(Path(__file__).resolve().parent))

from fastapi.testclient import TestClient  # noqa: E402
from auth.security import hash_password  # noqa: E402
from database import Base, SessionLocal, engine, run_migrations  # noqa: E402
from main import app  # noqa: E402
from models import (  # noqa: E402
    MesJobPackage, MesProductionJob, MesProductTemplate, ProductionProject,
    ProductionProjectLine, ProjectReleaseSnapshot, User,
)
from services.package_traceability import (  # noqa: E402
    _effective_passport_status, _timeline_event, _upsert_timeline_event,
)


def _login(client: TestClient, username: str) -> dict[str, str]:
    response = client.post("/auth/login", json={"username": username, "password": "11111111"})
    assert response.status_code == 200, response.text
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def test_passport_logistics_status_and_timeline_upsert_are_authoritative() -> None:
    assert _effective_passport_status("placed", "accepted") == "accepted"
    assert _effective_passport_status("placed", "planned") == "in_warehouse"
    timeline = [{"stage_key": "loading", "completed_at": None, "result": None}]
    event = _timeline_event(
        "loading", "Loading", operator="trace-admin",
        started_at=None, completed_at=datetime(2026, 1, 1, 12),
        result="loaded",
    )
    _upsert_timeline_event(timeline, event)
    _upsert_timeline_event(timeline, event)
    assert len(timeline) == 1
    assert timeline[0]["operator"] == "trace-admin"
    assert timeline[0]["result"] == "loaded"


def test_product_unit_passports_are_idempotent_and_traceable() -> None:
    Base.metadata.create_all(engine); run_migrations()
    db = SessionLocal()
    try:
        admin = User(username="trace-admin", password_hash=hash_password("11111111"), role="admin", department="Admin")
        outsider = User(username="trace-outsider", password_hash=hash_password("11111111"), role="operator", department="CRM")
        template = MesProductTemplate(code="TRACE-Z1", name="Traceable Z1", created_by="trace-admin")
        project = ProductionProject(project_code="TRACE-P", project_name="Trace project", customer_name_snapshot="Customer", destination_city="Toshkent", site_name="Atlas", full_address="Atlas", status="released", created_by="trace-admin")
        db.add_all([admin, outsider, template, project]); db.flush()
        line = ProductionProjectLine(project_id=project.id, product_id=template.id, quantity=2, status="released")
        release = ProjectReleaseSnapshot(project_id=project.id, revision=1, source_project_version=1, idempotency_key="trace-release", checksum="a" * 64, released_by="trace-admin")
        db.add_all([line, release]); db.flush()
        job = MesProductionJob(job_number="TRACE-JOB", customer_name="Customer", template_id=template.id, quantity=2, status="completed", created_by="trace-admin", project_id=project.id, project_line_id=line.id, project_release_snapshot_id=release.id)
        db.add(job); db.flush()
        package = MesJobPackage(job_id=job.id, package_number="TRACE-PKG", quantity=2, status="received", project_id=project.id, project_line_id=line.id, project_release_snapshot_id=release.id)
        db.add(package); db.commit(); package_id = package.id

        client = TestClient(app)
        admin_h = _login(client, "trace-admin")
        outsider_h = _login(client, "trace-outsider")
        assert client.get("/traceability/products", headers=outsider_h).status_code == 403
        assert client.post(f"/traceability/packages/{package_id}/passports", headers=outsider_h).status_code == 403
        first = client.post(f"/traceability/packages/{package_id}/passports", headers=admin_h)
        assert first.status_code == 200, first.text
        rows = first.json()["items"]
        assert len(rows) == 2 and len({row["serial_number"] for row in rows}) == 2
        assert all(row["qr_data"].startswith("https://trace.test/track/product/") for row in rows)
        second = client.post(f"/traceability/packages/{package_id}/passports", headers=admin_h)
        assert second.status_code == 200 and len(second.json()["items"]) == 2
        batch = client.post("/traceability/passports/batch", headers=admin_h, json={"package_ids": [package_id, package_id]})
        assert batch.status_code == 200 and len(batch.json()["items"]) == 2
        serial = rows[0]["serial_number"]
        detail = client.get(f"/traceability/products/{serial}", headers=admin_h)
        assert detail.status_code == 200 and detail.json()["project"]["code"] == "TRACE-P"
        token = detail.json()["qr_token"]
        public = client.get(f"/track/product/{token}")
        assert public.status_code == 200 and public.json()["serial_number"] == serial
        assert public.json().get("qr_token") is None
        assert all(row.get("operator") is None for row in public.json().get("timeline", []))
        label = client.get(f"/traceability/products/{serial}/label.png", headers=admin_h)
        assert label.status_code == 200 and label.headers["content-type"].startswith("image/png") and label.content.startswith(b"\x89PNG")
        for size in ("60x60", "a4"):
            variant = client.get(f"/traceability/products/{serial}/label.png?size={size}", headers=admin_h)
            assert variant.status_code == 200 and variant.content.startswith(b"\x89PNG")
        pdf = client.get(f"/traceability/products/{serial}/label.pdf", headers=admin_h)
        assert pdf.status_code == 200 and pdf.headers["content-type"].startswith("application/pdf") and pdf.content.startswith(b"%PDF-")
        passport_pdf = client.get(f"/traceability/products/{serial}/passport.pdf?language=uz", headers=admin_h)
        assert passport_pdf.status_code == 200 and passport_pdf.content.startswith(b"%PDF-")
        from pypdf import PdfReader
        assert len(PdfReader(__import__("io").BytesIO(passport_pdf.content)).pages) >= 1
        passport_docx = client.get(f"/traceability/products/{serial}/passport.docx?language=ru", headers=admin_h)
        assert passport_docx.status_code == 200 and passport_docx.content.startswith(b"PK")
        from docx import Document
        assert Document(__import__("io").BytesIO(passport_docx.content)).tables
        passport_xlsx = client.get(f"/traceability/products/{serial}/passport.xlsx?language=uz", headers=admin_h)
        assert passport_xlsx.status_code == 200 and passport_xlsx.content.startswith(b"PK")
        from openpyxl import load_workbook
        workbook = load_workbook(__import__("io").BytesIO(passport_xlsx.content))
        assert workbook.sheetnames == ["Mahsulot pasporti", "Ishlab chiqarish tarixi", "QC-Qayta ishlash", "Logistika"]
        batch_xlsx = client.post("/traceability/passports/batch/export.xlsx", headers=admin_h, json={"serial_numbers": [row["serial_number"] for row in rows], "language": "ru"})
        assert batch_xlsx.status_code == 200 and batch_xlsx.content.startswith(b"PK")
        artifact_dir = os.getenv("TRACEABILITY_ARTIFACT_DIR", "").strip()
        if artifact_dir:
            target = Path(artifact_dir); target.mkdir(parents=True, exist_ok=True)
            (target / f"{serial}-passport-uz.pdf").write_bytes(passport_pdf.content)
            (target / f"{serial}-passport-ru.docx").write_bytes(passport_docx.content)
            (target / f"{serial}-passport-uz.xlsx").write_bytes(passport_xlsx.content)
            (target / "product-passports-batch-ru.xlsx").write_bytes(batch_xlsx.content)
        sheet = client.post("/traceability/passports/batch/labels.pdf", headers=admin_h, json={"serial_numbers": [row["serial_number"] for row in rows], "size": "60x60"})
        assert sheet.status_code == 200 and sheet.headers["content-type"].startswith("application/pdf") and sheet.content.startswith(b"%PDF-")
    finally:
        db.close()
