"""Professional PDF, Word and Excel exports for canonical product passports."""

from __future__ import annotations

import io
from datetime import datetime
from typing import Any

from services.corporate_documents import (
    AMBER, BORDER, GREEN, NAVY, PALE, RED, SLATE, WHITE,
    corporate_branding, duration_label, format_datetime, language, qr_png,
    result_color, stage_label, status_label,
)


TEXT = {
    "uz": {
        "title": "MAHSULOT PASPORTI", "document": "Hujjat raqami", "date": "Sana",
        "product": "Mahsulot ma’lumotlari", "history": "Ishlab chiqarish tarixi",
        "quality": "Sifat nazorati va qayta ishlash", "packaging": "Qadoqlash",
        "logistics": "Logistika", "serial": "Seriya raqami", "code": "Mahsulot kodi",
        "name": "Mahsulot nomi", "project": "Loyiha", "quantity": "Birlik miqdori",
        "status": "Holat", "stage": "Bosqich", "operator": "Operator",
        "started": "Boshlangan", "completed": "Yakunlangan", "duration": "Davomiyligi",
        "result": "Natija", "unknown": "Noma’lum", "no_data": "Ma’lumot mavjud emas",
        "package": "Qadoq", "shipment": "Reys / jo‘natma", "destination": "Manzil",
        "vehicle": "Transport", "driver": "Haydovchi", "loaded": "Yuklangan",
        "delivered": "Yetkazilgan", "accepted": "Qabul qilingan", "issue": "Muammo",
        "detected": "Aniqlangan", "resolved": "Hal qilingan", "notes": "Izoh",
        "prepared": "Tayyorladi", "approved": "Tasdiqladi", "signature": "Imzo",
        "qr_help": "QR-kodni skanerlab mahsulot pasportini oching",
        "sheet_passport": "Mahsulot pasporti", "sheet_history": "Ishlab chiqarish tarixi",
        "sheet_quality": "QC-Qayta ishlash", "sheet_logistics": "Logistika",
    },
    "ru": {
        "title": "ПАСПОРТ ПРОДУКЦИИ", "document": "Номер документа", "date": "Дата",
        "product": "Сведения о продукции", "history": "История производства",
        "quality": "Контроль качества и доработка", "packaging": "Упаковка",
        "logistics": "Логистика", "serial": "Серийный номер", "code": "Код продукции",
        "name": "Наименование продукции", "project": "Проект", "quantity": "Количество единицы",
        "status": "Статус", "stage": "Этап", "operator": "Оператор",
        "started": "Начато", "completed": "Завершено", "duration": "Длительность",
        "result": "Результат", "unknown": "Неизвестно", "no_data": "Данные отсутствуют",
        "package": "Упаковка", "shipment": "Рейс / отправка", "destination": "Назначение",
        "vehicle": "Транспорт", "driver": "Водитель", "loaded": "Погружено",
        "delivered": "Доставлено", "accepted": "Принято", "issue": "Проблема",
        "detected": "Обнаружено", "resolved": "Решено", "notes": "Примечание",
        "prepared": "Подготовил", "approved": "Утвердил", "signature": "Подпись",
        "qr_help": "Отсканируйте QR-код, чтобы открыть паспорт продукции",
        "sheet_passport": "Паспорт продукции", "sheet_history": "История производства",
        "sheet_quality": "QC-Доработка", "sheet_logistics": "Логистика",
    },
}


def _t(lang: str, key: str) -> str:
    return TEXT[language(lang)][key]


def _project(data: dict) -> str:
    return " — ".join(filter(None, [data.get("project", {}).get("code"), data.get("project", {}).get("name")])) or "—"


def build_passport_pdf(db, data: dict[str, Any], *, lang: str = "uz") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    lang = language(lang); brand = corporate_branding(db); buf = io.BytesIO()
    regular_font, bold_font = "Helvetica", "Helvetica-Bold"
    try:
        pdfmetrics.registerFont(TTFont("VelkoreSans", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"))
        pdfmetrics.registerFont(TTFont("VelkoreSans-Bold", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"))
        regular_font, bold_font = "VelkoreSans", "VelkoreSans-Bold"
    except (OSError, IOError):
        pass
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=16*mm, rightMargin=16*mm, topMargin=17*mm, bottomMargin=17*mm,
                            title=f"{_t(lang, 'title')} {data['serial_number']}", author=brand["name"])
    styles = getSampleStyleSheet()
    body = ParagraphStyle("VelBody", parent=styles["BodyText"], fontName=regular_font, fontSize=8.2, leading=10.5, textColor=colors.HexColor("#0F172A"))
    small = ParagraphStyle("VelSmall", parent=body, fontSize=7, leading=8.5, textColor=colors.HexColor("#475569"))
    title = ParagraphStyle("VelTitle", parent=styles["Title"], fontName=bold_font, fontSize=16, leading=19, textColor=colors.HexColor("#FFFFFF"), alignment=TA_LEFT)
    section = ParagraphStyle("VelSection", parent=body, fontName=bold_font, fontSize=9.5, textColor=colors.HexColor("#FFFFFF"))
    table_style = TableStyle([("GRID", (0,0), (-1,-1), .4, colors.HexColor(f"#{BORDER}")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("LEFTPADDING", (0,0), (-1,-1), 5), ("RIGHTPADDING", (0,0), (-1,-1), 5), ("TOPPADDING", (0,0), (-1,-1), 4), ("BOTTOMPADDING", (0,0), (-1,-1), 4)])
    def P(value, style=body): return Paragraph(str(value if value not in (None, "") else "—"), style)
    def section_bar(text):
        table = Table([[Paragraph(text, section)]], colWidths=[178*mm]); table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,-1), colors.HexColor(f"#{NAVY}")), ("LEFTPADDING", (0,0), (-1,-1), 7), ("TOPPADDING", (0,0), (-1,-1), 5), ("BOTTOMPADDING", (0,0), (-1,-1), 5)])); return table
    header_left = [Paragraph(brand["name"].upper(), title), Paragraph(brand["tagline"], ParagraphStyle("tag", parent=small, textColor=colors.white))]
    header = Table([[header_left, P(f"{_t(lang, 'document')}: {data['serial_number']}<br/>{_t(lang, 'date')}: {format_datetime(datetime.utcnow(), lang)}", ParagraphStyle("meta", parent=small, textColor=colors.white, alignment=TA_LEFT))]], colWidths=[108*mm, 70*mm])
    header.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,-1),colors.HexColor(f"#{NAVY}")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),9),("RIGHTPADDING",(0,0),(-1,-1),9),("TOPPADDING",(0,0),(-1,-1),10),("BOTTOMPADDING",(0,0),(-1,-1),10)]))
    story = [header, Spacer(1, 7*mm), section_bar(_t(lang,"product")), Spacer(1, 2*mm)]
    info = [[P(_t(lang,"serial"), small), P(data["serial_number"]), P(_t(lang,"status"), small), P(status_label(data.get("status"), lang))], [P(_t(lang,"code"),small),P(data.get("product_code")),P(_t(lang,"name"),small),P(data.get("product_name"))], [P(_t(lang,"project"),small),P(_project(data)),P(_t(lang,"quantity"),small),P(data.get("unit_quantity"))]]
    info_table=Table(info,colWidths=[28*mm,57*mm,28*mm,65*mm]); info_table.setStyle(table_style); info_table.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),colors.HexColor(f"#{PALE}")),("BACKGROUND",(2,0),(2,-1),colors.HexColor(f"#{PALE}")),("FONTNAME",(0,0),(0,-1),bold_font),("FONTNAME",(2,0),(2,-1),bold_font)]))
    story += [info_table, Spacer(1,5*mm), section_bar(_t(lang,"history")), Spacer(1,2*mm)]
    history_headers=[_t(lang,x) for x in ("stage","operator","started","completed","duration","result")]
    history=[[P(x,small) for x in history_headers]]
    for row in data.get("timeline") or []:
        history.append([P(stage_label(row.get("stage_name"),lang)),P(row.get("operator") or _t(lang,"unknown")),P(format_datetime(row.get("started_at"),lang)),P(format_datetime(row.get("completed_at"),lang)),P(duration_label(row.get("duration_seconds"),lang)),P(status_label(row.get("result"),lang))])
    if len(history)==1: history.append([P(_t(lang,"no_data"))]+[P("—")]*5)
    hist=Table(history,colWidths=[29*mm,31*mm,31*mm,31*mm,22*mm,34*mm],repeatRows=1); hist.setStyle(table_style); hist.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor(f"#{PALE}")),("FONTNAME",(0,0),(-1,0),bold_font)]))
    story += [hist,Spacer(1,5*mm),section_bar(_t(lang,"quality")),Spacer(1,2*mm)]
    quality=[[P(x,small) for x in (_t(lang,"issue"),_t(lang,"status"),_t(lang,"operator"),_t(lang,"detected"),_t(lang,"resolved"),_t(lang,"notes"))]]
    for issue in data.get("quality",{}).get("issues") or []:
        quality.append([P(issue.get("stage") or issue.get("issue_type")),P(status_label(issue.get("status"),lang)),P(issue.get("resolved_by") or issue.get("detected_by") or _t(lang,"unknown")),P(format_datetime(issue.get("detected_at"),lang)),P(format_datetime(issue.get("resolved_at"),lang)),P(issue.get("notes"))])
    if len(quality)==1: quality.append([P(_t(lang,"no_data"))]+[P("—")]*5)
    qt=Table(quality,colWidths=[28*mm,25*mm,31*mm,29*mm,29*mm,36*mm],repeatRows=1); qt.setStyle(table_style); qt.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor(f"#{PALE}")),("FONTNAME",(0,0),(-1,0),bold_font)]))
    package=data.get("package") or {}; shipment=data.get("shipment") or {}
    logistics=[[P(_t(lang,"package"),small),P(package.get("number")),P(_t(lang,"status"),small),P(status_label(package.get("status"),lang))],[P(_t(lang,"shipment"),small),P(shipment.get("trip_number") or shipment.get("dispatch_number")),P(_t(lang,"destination"),small),P(shipment.get("destination"))],[P(_t(lang,"vehicle"),small),P(shipment.get("vehicle")),P(_t(lang,"driver"),small),P(shipment.get("driver") or _t(lang,"unknown"))],[P(_t(lang,"loaded"),small),P(format_datetime(shipment.get("loaded_at"),lang)),P(_t(lang,"delivered"),small),P(format_datetime(shipment.get("delivered_at"),lang))]]
    lt=Table(logistics,colWidths=[28*mm,57*mm,28*mm,65*mm]);lt.setStyle(table_style);lt.setStyle(TableStyle([("BACKGROUND",(0,0),(0,-1),colors.HexColor(f"#{PALE}")),("BACKGROUND",(2,0),(2,-1),colors.HexColor(f"#{PALE}")),("FONTNAME",(0,0),(0,-1),bold_font),("FONTNAME",(2,0),(2,-1),bold_font)]))
    qr=Image(io.BytesIO(qr_png(data.get("qr_data") or "",pixels=360)),width=36*mm,height=36*mm)
    qr_panel=Table([[qr,[P(data["serial_number"],ParagraphStyle("qrserial",parent=body,fontName=bold_font,fontSize=10)),P(_t(lang,"qr_help"),small)]]],colWidths=[42*mm,136*mm]);qr_panel.setStyle(TableStyle([("BOX",(0,0),(-1,-1),.5,colors.HexColor(f"#{BORDER}")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("LEFTPADDING",(0,0),(-1,-1),6),("TOPPADDING",(0,0),(-1,-1),6),("BOTTOMPADDING",(0,0),(-1,-1),6)]))
    story += [qt,Spacer(1,5*mm),section_bar(_t(lang,"logistics")),Spacer(1,2*mm),lt,Spacer(1,5*mm),KeepTogether(qr_panel)]
    def page(canvas, document):
        canvas.saveState(); canvas.setStrokeColor(colors.HexColor(f"#{BORDER}")); canvas.line(16*mm,12*mm,194*mm,12*mm); canvas.setFillColor(colors.HexColor(f"#{SLATE}")); canvas.setFont(regular_font,7); canvas.drawString(16*mm,7.5*mm,f"{brand['name']} · {data['serial_number']}"); canvas.drawRightString(194*mm,7.5*mm,f"{document.page}"); canvas.restoreState()
    doc.build(story,onFirstPage=page,onLaterPages=page); return buf.getvalue()


def build_passport_docx(db, data: dict[str, Any], *, lang: str = "uz") -> bytes:
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    lang=language(lang); brand=corporate_branding(db); doc=Document(); section=doc.sections[0]
    section.page_height=Cm(29.7); section.page_width=Cm(21); section.top_margin=Cm(1.6); section.bottom_margin=Cm(1.6); section.left_margin=Cm(1.7); section.right_margin=Cm(1.7)
    styles=doc.styles; styles["Normal"].font.name="Arial"; styles["Normal"].font.size=Pt(9); styles["Normal"].paragraph_format.space_after=Pt(2)
    def shade(cell, fill): cell._tc.get_or_add_tcPr().append(OxmlElement("w:shd")); cell._tc.tcPr[-1].set(qn("w:fill"),fill)
    def set_cell(cell,text,bold=False,color=None,size=8):
        cell.text=""; p=cell.paragraphs[0]; r=p.add_run(str(text if text not in (None,"") else "—")); r.bold=bold; r.font.name="Arial"; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color or "0F172A"); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
    header=section.header; table=header.add_table(rows=1,cols=2,width=Cm(17.6)); table.columns[0].width=Cm(11); table.columns[1].width=Cm(6.6)
    set_cell(table.cell(0,0),f"{brand['name'].upper()}\n{brand['tagline']}",True,WHITE,13); set_cell(table.cell(0,1),f"{_t(lang,'document')}: {data['serial_number']}\n{_t(lang,'date')}: {format_datetime(datetime.utcnow(),lang)}",False,WHITE,8); shade(table.cell(0,0),NAVY); shade(table.cell(0,1),NAVY)
    title=doc.add_paragraph();title.alignment=WD_ALIGN_PARAGRAPH.CENTER;r=title.add_run(_t(lang,"title"));r.bold=True;r.font.size=Pt(16);r.font.color.rgb=RGBColor.from_string(NAVY)
    def heading(text):
        t=doc.add_table(rows=1,cols=1);t.alignment=WD_TABLE_ALIGNMENT.CENTER;set_cell(t.cell(0,0),text,True,WHITE,9);shade(t.cell(0,0),NAVY);doc.add_paragraph().paragraph_format.space_after=Pt(0)
    def grid(rows,widths=None,header=False):
        table=doc.add_table(rows=0,cols=len(rows[0]));table.style="Table Grid";table.alignment=WD_TABLE_ALIGNMENT.CENTER
        for ri,row in enumerate(rows):
            cells=table.add_row().cells
            for ci,value in enumerate(row): set_cell(cells[ci],value,bold=header and ri==0,size=8); shade(cells[ci],PALE if header and ri==0 else WHITE)
        if widths:
            for row in table.rows:
                for idx,width in enumerate(widths): row.cells[idx].width=Cm(width)
        return table
    heading(_t(lang,"product")); grid([[_t(lang,"serial"),data["serial_number"],_t(lang,"status"),status_label(data.get("status"),lang)],[_t(lang,"code"),data.get("product_code"),_t(lang,"name"),data.get("product_name")],[_t(lang,"project"),_project(data),_t(lang,"quantity"),data.get("unit_quantity")]], [3,5.8,3,5.8])
    heading(_t(lang,"history")); rows=[[_t(lang,x) for x in ("stage","operator","started","completed","duration","result")]]
    for row in data.get("timeline") or []: rows.append([stage_label(row.get("stage_name"),lang),row.get("operator") or _t(lang,"unknown"),format_datetime(row.get("started_at"),lang),format_datetime(row.get("completed_at"),lang),duration_label(row.get("duration_seconds"),lang),status_label(row.get("result"),lang)])
    if len(rows)==1: rows.append([_t(lang,"no_data"),"—","—","—","—","—"])
    grid(rows,[2.7,3,3.1,3.1,2.2,3.5],True)
    heading(_t(lang,"quality")); rows=[[_t(lang,x) for x in ("issue","status","operator","detected","resolved","notes")]]
    for issue in data.get("quality",{}).get("issues") or []: rows.append([issue.get("stage") or issue.get("issue_type"),status_label(issue.get("status"),lang),issue.get("resolved_by") or issue.get("detected_by") or _t(lang,"unknown"),format_datetime(issue.get("detected_at"),lang),format_datetime(issue.get("resolved_at"),lang),issue.get("notes")])
    if len(rows)==1: rows.append([_t(lang,"no_data"),"—","—","—","—","—"])
    grid(rows,[2.6,2.4,3,3,3,3.6],True)
    heading(_t(lang,"logistics")); package=data.get("package") or {};shipment=data.get("shipment") or {};grid([[_t(lang,"package"),package.get("number"),_t(lang,"status"),status_label(package.get("status"),lang)],[_t(lang,"shipment"),shipment.get("trip_number") or shipment.get("dispatch_number"),_t(lang,"destination"),shipment.get("destination")],[_t(lang,"vehicle"),shipment.get("vehicle"),_t(lang,"driver"),shipment.get("driver") or _t(lang,"unknown")],[_t(lang,"loaded"),format_datetime(shipment.get("loaded_at"),lang),_t(lang,"delivered"),format_datetime(shipment.get("delivered_at"),lang)]],[3,5.8,3,5.8])
    qr_table=doc.add_table(rows=1,cols=2);qr_table.style="Table Grid"; qr_table.cell(0,0).paragraphs[0].add_run().add_picture(io.BytesIO(qr_png(data.get("qr_data") or "")),width=Cm(3.6));set_cell(qr_table.cell(0,1),f"{data['serial_number']}\n{_t(lang,'qr_help')}",True,NAVY,9)
    doc.add_paragraph(); sig=doc.add_table(rows=2,cols=3);sig.style="Table Grid";set_cell(sig.cell(0,0),_t(lang,"prepared"),True);set_cell(sig.cell(0,1),_t(lang,"approved"),True);set_cell(sig.cell(0,2),_t(lang,"signature"),True);[shade(c,PALE) for c in sig.rows[0].cells];[set_cell(c,"\n") for c in sig.rows[1].cells]
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;run=footer.add_run(f"{brand['name']} · {data['serial_number']} · ");run.font.size=Pt(8);fld=OxmlElement("w:fldSimple");fld.set(qn("w:instr"),"PAGE");footer._p.append(fld)
    out=io.BytesIO();doc.save(out);return out.getvalue()


def _style_workbook(workbook, sheet, *, title: str, subtitle: str, columns: int):
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    sheet.merge_cells(start_row=1,start_column=1,end_row=1,end_column=columns);sheet["A1"]=title;sheet["A1"].font=Font(name="Arial",size=16,bold=True,color=WHITE);sheet["A1"].fill=PatternFill("solid",fgColor=NAVY);sheet["A1"].alignment=Alignment(horizontal="center",vertical="center");sheet.row_dimensions[1].height=30
    sheet.merge_cells(start_row=2,start_column=1,end_row=2,end_column=columns);sheet["A2"]=subtitle;sheet["A2"].font=Font(name="Arial",size=9,color=SLATE);sheet["A2"].alignment=Alignment(horizontal="center")
    sheet.freeze_panes="A4";sheet.sheet_view.showGridLines=False;sheet.page_setup.orientation="landscape";sheet.page_setup.paperSize=sheet.PAPERSIZE_A4;sheet.page_setup.fitToWidth=1;sheet.page_margins.left=.3;sheet.page_margins.right=.3;sheet.page_margins.top=.5;sheet.page_margins.bottom=.5
    sheet.oddFooter.center.text="Velkore ERP · Page &P of &N"
    return Border(*( [Side(style="thin",color=BORDER)] * 4 ))


def build_passport_xlsx(db, data: dict[str, Any], *, lang: str = "uz") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    lang=language(lang);brand=corporate_branding(db);wb=Workbook();wb.remove(wb.active)
    sheets=[(_t(lang,"sheet_passport"),4),(_t(lang,"sheet_history"),6),(_t(lang,"sheet_quality"),6),(_t(lang,"sheet_logistics"),4)]
    for name,cols in sheets: ws=wb.create_sheet(name[:31]);_style_workbook(wb,ws,title=_t(lang,"title"),subtitle=f"{brand['name']} · {data['serial_number']}",columns=cols)
    def header(ws,row,values):
        for col,value in enumerate(values,1): c=ws.cell(row,col,value);c.font=Font(name="Arial",bold=True,color=WHITE);c.fill=PatternFill("solid",fgColor=NAVY);c.alignment=Alignment(wrap_text=True,vertical="center")
    def body(ws,row,values):
        for col,value in enumerate(values,1): c=ws.cell(row,col,value if value not in (None,"") else "—");c.font=Font(name="Arial",size=9);c.alignment=Alignment(wrap_text=True,vertical="top")
    ws=wb.worksheets[0];header(ws,4,[_t(lang,"serial"),_t(lang,"code"),_t(lang,"name"),_t(lang,"status")]);body(ws,5,[data["serial_number"],data.get("product_code"),data.get("product_name"),status_label(data.get("status"),lang)]);header(ws,7,[_t(lang,"project"),_t(lang,"package"),_t(lang,"quantity"),_t(lang,"date")]);body(ws,8,[_project(data),data.get("package",{}).get("number"),data.get("unit_quantity"),format_datetime(data.get("created_at"),lang)]);ws.auto_filter.ref="A4:D5"
    ws=wb.worksheets[1];header(ws,4,[_t(lang,x) for x in ("stage","operator","started","completed","duration","result")]);r=5
    for item in data.get("timeline") or []: body(ws,r,[stage_label(item.get("stage_name"),lang),item.get("operator") or _t(lang,"unknown"),item.get("started_at"),item.get("completed_at"),duration_label(item.get("duration_seconds"),lang),status_label(item.get("result"),lang)]);r+=1
    ws.auto_filter.ref=f"A4:F{max(4,r-1)}"
    ws=wb.worksheets[2];header(ws,4,[_t(lang,x) for x in ("issue","status","operator","detected","resolved","notes")]);r=5
    for item in data.get("quality",{}).get("issues") or []: body(ws,r,[item.get("stage") or item.get("issue_type"),status_label(item.get("status"),lang),item.get("resolved_by") or item.get("detected_by") or _t(lang,"unknown"),item.get("detected_at"),item.get("resolved_at"),item.get("notes")]);r+=1
    ws.auto_filter.ref=f"A4:F{max(4,r-1)}"
    ws=wb.worksheets[3];shipment=data.get("shipment") or {};package=data.get("package") or {};header(ws,4,[_t(lang,"package"),_t(lang,"shipment"),_t(lang,"vehicle"),_t(lang,"driver")]);body(ws,5,[package.get("number"),shipment.get("trip_number") or shipment.get("dispatch_number"),shipment.get("vehicle"),shipment.get("driver") or _t(lang,"unknown")]);header(ws,7,[_t(lang,"destination"),_t(lang,"loaded"),_t(lang,"delivered"),_t(lang,"accepted")]);body(ws,8,[shipment.get("destination"),shipment.get("loaded_at"),shipment.get("delivered_at"),shipment.get("accepted_at")]);ws.auto_filter.ref="A4:D5"
    for ws in wb.worksheets:
        for width,col in zip((24,24,30,24,22,34),"ABCDEF"): ws.column_dimensions[col].width=width
        ws.print_title_rows="1:4"
        for row in ws.iter_rows(min_row=4):
            for cell in row: cell.border = __import__("openpyxl").styles.Border(left=__import__("openpyxl").styles.Side(style="thin",color=BORDER),right=__import__("openpyxl").styles.Side(style="thin",color=BORDER),top=__import__("openpyxl").styles.Side(style="thin",color=BORDER),bottom=__import__("openpyxl").styles.Side(style="thin",color=BORDER))
    out=io.BytesIO();wb.save(out);return out.getvalue()


def build_passports_batch_xlsx(db, passports: list[dict[str, Any]], *, lang: str = "uz") -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    lang=language(lang);brand=corporate_branding(db);wb=Workbook();ws=wb.active;ws.title="Passports" if lang=="ru" else "Pasportlar";_style_workbook(wb,ws,title=_t(lang,"title"),subtitle=f"{brand['name']} · {len(passports)}",columns=10)
    headers=[_t(lang,"serial"),_t(lang,"code"),_t(lang,"name"),_t(lang,"project"),_t(lang,"package"),_t(lang,"status"),_t(lang,"shipment"),_t(lang,"destination"),_t(lang,"delivered"),"QR URL"]
    for c,value in enumerate(headers,1): cell=ws.cell(4,c,value);cell.font=Font(bold=True,color=WHITE);cell.fill=PatternFill("solid",fgColor=NAVY);cell.alignment=Alignment(wrap_text=True)
    for r,data in enumerate(passports,5):
        shipment=data.get("shipment") or {};values=[data.get("serial_number"),data.get("product_code"),data.get("product_name"),_project(data),data.get("package",{}).get("number"),status_label(data.get("status"),lang),shipment.get("trip_number") or shipment.get("dispatch_number"),shipment.get("destination"),shipment.get("delivered_at"),data.get("qr_data")]
        for c,value in enumerate(values,1):ws.cell(r,c,value if value not in (None,"") else "—").alignment=Alignment(wrap_text=True,vertical="top")
    ws.auto_filter.ref=f"A4:J{max(4,len(passports)+4)}";ws.freeze_panes="A5";ws.print_title_rows="1:4"
    for col,width in zip("ABCDEFGHIJ",(22,18,28,28,18,20,22,30,22,45)):ws.column_dimensions[col].width=width
    out=io.BytesIO();wb.save(out);return out.getvalue()
